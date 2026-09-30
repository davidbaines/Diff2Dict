"""Excel workbook and CSV report writers."""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .mine import Result, _show_ctx

MAX_WIDTH = 60


def _clean(value):
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("�", value)
    return value


def _sheet(wb: Workbook, title: str, header: list[str], rows) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    widths = [len(h) for h in header]
    for row in rows:
        ws.append([_clean(v) for v in row])
        for cell in ws[ws.max_row]:
            # Text such as "=" or "-" must stay text, not become a formula.
            if isinstance(cell.value, str):
                cell.data_type = "s"
        for k, v in enumerate(row):
            widths[k] = min(MAX_WIDTH, max(widths[k], len(str(v))))
    for k, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(k)].width = w + 2
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def _p(value: float) -> float:
    return round(value, 3)


def _lines(numbers: list[int]) -> str:
    return ", ".join(map(str, numbers))


def write_workbook(path: str | Path, result: Result, summary: dict[str, object]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    in_map = {m.text for m in result.map_rules}

    _sheet(wb, "CharPairs",
           ["source", "target", "left_ctx", "right_ctx", "count", "p_fwd", "p_bwd",
            "direction", "in_map", "examples"],
           ([ "".join(r.x), "".join(r.y), _show_ctx(r.left), _show_ctx(r.right), r.count,
              _p(r.p_fwd), _p(r.p_bwd),
              "rejected" if r.rejected else (r.direction or "unreliable"),
              "yes" if r.text() in in_map else "no", "; ".join(r.examples)]
            for r in result.char_rules))

    genuine = [w for w in result.word_pairs if w.genuine]
    _sheet(wb, "WordPairs",
           ["source", "target", "kind", "count", "n_source", "n_target", "p_fwd", "p_bwd",
            "direction", "in_map", "example_lines"],
           ([w.source, w.target, w.kind, w.count, w.n_source, w.n_target, _p(w.p_fwd),
             _p(w.p_bwd), w.direction, "yes" if w.direction else "no", _lines(w.lines)]
            for w in genuine))

    lexicon = [w for w in result.word_pairs if w.kind not in ("insertion", "deletion")]
    _sheet(wb, "Lexicon", ["source", "target", "kind", "count", "explained_by"],
           ([w.source, w.target, "regular" if w.kind == "regular" else "genuine",
             w.count, "; ".join(dict.fromkeys(w.rules))] for w in lexicon))

    _sheet(wb, "PunctPairs",
           ["source", "target", "count", "p_fwd", "p_bwd", "direction", "example_lines"],
           ([p.source, p.target, p.count, _p(p.p_fwd), _p(p.p_bwd), p.direction or "unreliable",
             _lines(p.lines)] for p in result.punct_pairs))

    _sheet(wb, "CaseOnly", ["source", "target", "count"],
           ([s, t, c] for (s, t), c in result.case_only.most_common()))

    _sheet(wb, "SkippedLines", ["line", "similarity", "source", "target"],
           ([line.number, _p(sim), line.src, line.tgt] for line, sim in result.skipped))

    _sheet(wb, "Summary", ["item", "value"],
           ([k, v if isinstance(v, (int, float)) else str(v)] for k, v in summary.items()))
    wb.move_sheet("Summary", offset=-(len(wb.sheetnames) - 1))
    wb.save(path)


def write_oneway_csv(path: str | Path, rows: list[list]) -> None:
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["rule", "kind", "direction", "p_fwd", "p_bwd", "reason"])
        w.writerows(rows)
