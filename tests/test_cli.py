import shutil

import openpyxl
import pytest

from diff2dict import report
from diff2dict.cli import main, select_files

from .conftest import SAMPLE, teckit_missing


def copy_sample(tmp_path):
    for name in ("source.txt", "target.txt"):
        shutil.copy(SAMPLE / name, tmp_path / name)
    return tmp_path / "source.txt", tmp_path / "target.txt"


def base_args(src, tgt, tmp_path):
    return [str(src), str(tgt), "--chars", str(tmp_path / "chars.csv"),
            "--out", str(tmp_path / "r.xlsx"), "--map", str(tmp_path / "r.map"),
            "--holdout", "0"]


def test_check_chars_without_file_generates_them_and_stops(tmp_path):
    src, tgt = copy_sample(tmp_path)
    chars = tmp_path / "chars.csv"
    code = main(base_args(src, tgt, tmp_path) + ["--check-chars"], ask=lambda _: "")
    assert code == 2
    assert chars.exists() and not (tmp_path / "r.xlsx").exists()


def test_check_chars_with_file_pauses_then_runs(tmp_path):
    src, tgt = copy_sample(tmp_path)
    main(base_args(src, tgt, tmp_path) + ["--check-chars"], ask=lambda _: "")  # writes chars.csv
    asked = []
    code = main(base_args(src, tgt, tmp_path) + ["--check-chars"], ask=lambda p: asked.append(p) or "")
    assert code == 0 and asked and (tmp_path / "r.xlsx").exists()


def test_default_run_generates_chars_and_completes(tmp_path):
    src, tgt = copy_sample(tmp_path)
    chars, out, map_path = tmp_path / "chars.csv", tmp_path / "r.xlsx", tmp_path / "r.map"
    assert main(base_args(src, tgt, tmp_path)) == 0
    assert chars.exists()
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Summary", "CharPairs", "WordPairs", "Lexicon", "PunctPairs",
                             "CaseOnly", "SkippedLines"]
    word_pairs = [r[:3] for r in wb["WordPairs"].iter_rows(min_row=2, values_only=True)]
    assert ("pin", "hunu", "substitution") in word_pairs
    skipped = [r[0] for r in wb["SkippedLines"].iter_rows(min_row=2, values_only=True)]
    assert skipped == [15]
    assert map_path.exists() and (tmp_path / "oneway_rules.csv").exists()
    oneway = (tmp_path / "oneway_rules.csv").read_text(encoding="utf-8-sig")
    assert "e → ∅ / o _" in oneway


def test_select_files_by_extension_and_all_text(tmp_path):
    (tmp_path / "a.txt").write_text("hello", encoding="utf-8")
    (tmp_path / "b.sfm").write_text("\\v 1 world", encoding="utf-8")
    (tmp_path / "c.bin").write_bytes(b"\x00\x01\x02binary")
    assert [p.name for p in select_files(tmp_path, "txt")] == ["a.txt"]
    assert [p.name for p in select_files(tmp_path, ".sfm")] == ["b.sfm"]
    # All-text markers take every text file but skip the binary one.
    for marker in ("*", ".*", "*.*"):
        assert [p.name for p in select_files(tmp_path, marker)] == ["a.txt", "b.sfm"]


def test_conversion_requires_ext_and_output(tmp_path):
    src, tgt = copy_sample(tmp_path)
    with pytest.raises(SystemExit):
        main(base_args(src, tgt, tmp_path) + ["--input-folder", str(tmp_path)])


@teckit_missing
def test_full_pipeline_compiles_and_converts_a_folder(tmp_path):
    src, tgt = copy_sample(tmp_path)
    in_folder, out_folder = tmp_path / "in", tmp_path / "out"
    in_folder.mkdir()
    (in_folder / "one.txt").write_text((SAMPLE / "source.txt").read_text(encoding="utf-8"),
                                       encoding="utf-8")
    code = main(base_args(src, tgt, tmp_path)
                + ["--input-folder", str(in_folder), "--input-ext", "txt",
                   "--output-folder", str(out_folder), "--tec", str(tmp_path / "r.tec")])
    assert code == 0
    assert (tmp_path / "r.tec").exists()
    converted = (out_folder / "one.txt").read_text(encoding="utf-8").splitlines()
    target = (SAMPLE / "target.txt").read_text(encoding="utf-8").splitlines()
    # Line 1 (kala -> kaalaa) is converted by the compiled map.
    assert converted[0] == target[0]


def test_workbook_keeps_formula_like_text_as_text(tmp_path):
    wb = openpyxl.Workbook()
    report._sheet(wb, "T", ["a"], [["=1+1"], ["-"], ["bad\x01char"]])
    wb.save(tmp_path / "t.xlsx")
    ws = openpyxl.load_workbook(tmp_path / "t.xlsx")["T"]
    assert [c.value for c in ws["A"]][1:] == ["=1+1", "-", "bad�char"]
    assert ws["A2"].data_type == "s"
